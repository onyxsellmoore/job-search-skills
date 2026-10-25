#!/usr/bin/env python3
# skill-version: 2026-10-15
"""One-file workspace for the free-plan job search skill.

Each chat starts with an empty sandbox, so everything the skill knows about
the user lives in ONE Excel workbook the user keeps: the tracker sheets
(jobs, contacts, daily_runs, written by tracker_ops.py) plus a hidden
`_profile` sheet holding the text files of the profile folder.

Usage:
    workbook.py info   --xlsx FILE
    workbook.py unpack --xlsx FILE --ws DIR [--force]
    workbook.py pack   --ws DIR --out FILE [--owner NAME]

unpack writes DIR/application_tracker.xlsx (a copy of FILE) and every stored
profile file under DIR/profile/. pack stores DIR/profile/**/*.json and *.md
into the `_profile` sheet of DIR/application_tracker.xlsx and saves the
result as FILE. PDFs are never stored (the base resume is re-rendered from
profile/resume.json when needed).

Every command prints one JSON object. Exit codes: 0 ok; 1 the file is not a
workbook of this skill (or DIR already holds a profile and --force was not
given); 2 unreadable input or bad arguments.
"""
import argparse
import json
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"ok": False, "error": "openpyxl not installed"}))
    sys.exit(2)

FORMAT = "job-search-lite"
FORMAT_VERSION = 1
SHEET = "_profile"
META = "__meta__"
CHUNK = 30000  # Excel caps a cell at 32767 characters
STORED_SUFFIXES = {".json", ".md"}
SKILL_VERSION = "2026-10-15"


def out(obj, code=0):
    print(json.dumps(obj, indent=2))
    sys.exit(code)


def count_rows(ws):
    return sum(
        1 for row in ws.iter_rows(min_row=2, values_only=True)
        if row and any(c is not None and str(c).strip() != "" for c in row)
    )


def last_run_date(wb):
    if "daily_runs" not in wb.sheetnames:
        return None
    best = None
    for row in wb["daily_runs"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        v = row[0]
        s = v.date().isoformat() if isinstance(v, datetime) else str(v)[:10]
        if best is None or s > best:
            best = s
    return best


def read_store(wb):
    """Return (meta dict or None, {relpath: text})."""
    if SHEET not in wb.sheetnames:
        return None, {}
    ws = wb[SHEET]
    meta, parts = None, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        path, part, text = row[0], row[1], row[2] if len(row) > 2 else ""
        text = "" if text is None else str(text)
        if path == META:
            try:
                meta = json.loads(text)
            except ValueError:
                meta = None
            continue
        parts.setdefault(str(path), []).append((int(part or 0), text))
    files = {p: "".join(t for _, t in sorted(chunks)) for p, chunks in parts.items()}
    return meta, files


def safe_rel(path):
    """A stored path must be a plain relative path under profile/."""
    p = PurePosixPath(path)
    if p.is_absolute() or ".." in p.parts or not p.parts or p.parts[0] != "profile":
        return None
    if p.suffix not in STORED_SUFFIXES:
        return None
    return p


def load(xlsx):
    try:
        return load_workbook(xlsx)
    except Exception as e:  # noqa: BLE001 - any read failure is a bad input
        out({"ok": False, "error": f"cannot read {xlsx}: {e}"}, 2)


def describe(wb, meta):
    return {
        "saved_at": (meta or {}).get("saved_at"),
        "owner": (meta or {}).get("owner", ""),
        "jobs": count_rows(wb["jobs"]) if "jobs" in wb.sheetnames else 0,
        "jobs_at_save": (meta or {}).get("jobs"),
        "last_run_date": last_run_date(wb),
    }


def cmd_info(a):
    wb = load(a.xlsx)
    meta, files = read_store(wb)
    if "jobs" not in wb.sheetnames or not meta or meta.get("format") != FORMAT:
        out({"ok": False, "error": "not a job-search-lite workbook (no jobs sheet or no _profile store)"}, 1)
    res = {"ok": True, "xlsx": a.xlsx, "files": sorted(files)}
    res.update(describe(wb, meta))
    out(res)


def cmd_unpack(a):
    wb = load(a.xlsx)
    meta, files = read_store(wb)
    if "jobs" not in wb.sheetnames or not meta or meta.get("format") != FORMAT:
        out({"ok": False, "error": "not a job-search-lite workbook (no jobs sheet or no _profile store)"}, 1)
    ws_dir = Path(a.ws)
    if (ws_dir / "profile").exists() and not a.force:
        out({"ok": False, "error": f"{ws_dir}/profile already exists; pass --force to replace it"}, 1)
    bad = [p for p in files if safe_rel(p) is None]
    if bad:
        out({"ok": False, "error": f"refusing unsafe stored paths: {bad}"}, 1)
    if (ws_dir / "profile").exists():
        shutil.rmtree(ws_dir / "profile")
    ws_dir.mkdir(parents=True, exist_ok=True)
    for rel, text in files.items():
        dest = ws_dir.joinpath(*PurePosixPath(rel).parts)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(text, encoding="utf-8")
    tracker = ws_dir / "application_tracker.xlsx"
    if Path(a.xlsx).resolve() != tracker.resolve():
        shutil.copyfile(a.xlsx, tracker)
    res = {"ok": True, "ws": str(ws_dir), "tracker": str(tracker), "files": sorted(files)}
    res.update(describe(wb, meta))
    out(res)


def cmd_pack(a):
    ws_dir = Path(a.ws)
    tracker = ws_dir / "application_tracker.xlsx"
    pdir = ws_dir / "profile"
    if not tracker.is_file():
        out({"ok": False, "error": f"missing {tracker}"}, 2)
    if not pdir.is_dir():
        out({"ok": False, "error": f"missing {pdir}"}, 2)
    wb = load(str(tracker))
    if "jobs" not in wb.sheetnames:
        out({"ok": False, "error": f"{tracker} has no jobs sheet"}, 2)
    old_meta, _ = read_store(wb)
    files = {}
    for f in sorted(pdir.rglob("*")):
        if f.is_file() and f.suffix in STORED_SUFFIXES:
            rel = PurePosixPath(*f.relative_to(ws_dir).parts)
            files[str(rel)] = f.read_text(encoding="utf-8")
    if "profile/preferences.json" not in files or "profile/resume.json" not in files:
        out({"ok": False, "error": "profile/preferences.json and profile/resume.json are required"}, 2)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append(["path", "part", "text"])
    owner = a.owner or (old_meta or {}).get("owner", "")
    jobs = count_rows(wb["jobs"])
    meta = {
        "format": FORMAT,
        "format_version": FORMAT_VERSION,
        "skill_version": SKILL_VERSION,
        "saved_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "owner": owner,
        "jobs": jobs,
    }
    ws.append([META, 0, json.dumps(meta)])
    for rel, text in files.items():
        chunks = [text[i:i + CHUNK] for i in range(0, len(text), CHUNK)] or [""]
        for i, c in enumerate(chunks):
            ws.append([rel, i, c])
    ws.sheet_state = "hidden"
    if wb.active is not None and wb.active.title == SHEET:
        wb.active = 0
    dest = Path(a.out)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    # Round-trip check: the saved file must give back exactly what was stored.
    meta2, files2 = read_store(load_workbook(dest))
    if files2 != files or not meta2 or meta2.get("format") != FORMAT:
        out({"ok": False, "error": "round-trip check failed; the saved workbook does not match"}, 2)
    out({"ok": True, "out": str(dest), "files": sorted(files), "jobs": jobs, "saved_at": meta["saved_at"]})


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("info")
    s.add_argument("--xlsx", required=True)
    s = sub.add_parser("unpack")
    s.add_argument("--xlsx", required=True)
    s.add_argument("--ws", required=True)
    s.add_argument("--force", action="store_true")
    s = sub.add_parser("pack")
    s.add_argument("--ws", required=True)
    s.add_argument("--out", required=True)
    s.add_argument("--owner")
    a = p.parse_args(argv)
    {"info": cmd_info, "unpack": cmd_unpack, "pack": cmd_pack}[a.cmd](a)


if __name__ == "__main__":
    main()
